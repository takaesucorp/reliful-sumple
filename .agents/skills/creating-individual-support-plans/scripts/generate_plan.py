#!/usr/bin/env python3
"""Create an editable A4 DOCX individual support plan from validated JSON."""

from __future__ import annotations

import argparse
import json
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


SKILL_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SCHEMA = SKILL_ROOT / "references" / "plan-schema.json"
FONT_NAME = "Yu Gothic"
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--input", type=Path, help="Input JSON matching plan-schema.json")
    source.add_argument("--blank", action="store_true", help="Create an anonymous blank template")
    parser.add_argument("--output", type=Path, required=True, help="Output .docx path")
    parser.add_argument("--schema", type=Path, default=DEFAULT_SCHEMA, help="JSON schema path")
    return parser.parse_args()


def load_json(path: Path) -> dict[str, Any]:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ValueError(f"ファイルが見つかりません: {path}") from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f"JSONの形式が正しくありません: {path} ({exc})") from exc


def resolve_ref(root: dict[str, Any], ref: str) -> dict[str, Any]:
    if not ref.startswith("#/"):
        raise ValueError(f"外部参照には対応していません: {ref}")
    node: Any = root
    for part in ref[2:].split("/"):
        node = node[part]
    return node


def validate(data: Any, schema: dict[str, Any], root: dict[str, Any], path: str = "$") -> list[str]:
    if "$ref" in schema:
        return validate(data, resolve_ref(root, schema["$ref"]), root, path)

    errors: list[str] = []
    expected = schema.get("type")
    type_map = {"object": dict, "array": list, "string": str}
    if expected in type_map and not isinstance(data, type_map[expected]):
        return [f"{path}: {expected} が必要です"]

    if expected == "object":
        for key in schema.get("required", []):
            if key not in data:
                errors.append(f"{path}.{key}: 必須項目がありません")
        props = schema.get("properties", {})
        if schema.get("additionalProperties") is False:
            for key in data:
                if key not in props:
                    errors.append(f"{path}.{key}: 未定義の項目です")
        for key, value in data.items():
            if key in props:
                errors.extend(validate(value, props[key], root, f"{path}.{key}"))
    elif expected == "array":
        minimum = schema.get("minItems", 0)
        if len(data) < minimum:
            errors.append(f"{path}: {minimum}件以上必要です")
        item_schema = schema.get("items", {})
        for index, item in enumerate(data):
            errors.extend(validate(item, item_schema, root, f"{path}[{index}]"))
    return errors


def blank_data() -> dict[str, Any]:
    blank_row = {
        "label": "",
        "monday": "",
        "tuesday": "",
        "wednesday": "",
        "thursday": "",
        "friday": "",
        "saturday": "",
    }
    blank_item = {
        "domains": [],
        "goal": "",
        "start_date": "",
        "end_date": "",
        "support_details": [""],
        "assigned_roles": [""],
    }
    return {
        "metadata": {
            "target_year": "",
            "target_month": "",
            "user_name": "",
            "created_date": "",
            "previous_created_date": "",
            "author": "",
            "manager_approval": "",
            "service_manager_approval": "",
        },
        "wishes": {"person": "", "family": ""},
        "goals": {
            "long_term": {"text": "", "start_date": "", "end_date": ""},
            "short_term": {"text": "", "start_date": "", "end_date": ""},
        },
        "weekly_schedule": {"legend": "", "rows": [dict(blank_row) for _ in range(3)]},
        "support_items": [dict(blank_item) for _ in range(3)],
        "notes": "",
        "agreement": {
            "explanation_date": "",
            "person_name": "",
            "family_name": "",
            "consent_date": "",
            "service_confirmation_date": "",
            "service_confirmation_name": "",
        },
        "facility": {
            "service_type": "就労継続支援B型",
            "name": "",
            "postal_code": "",
            "address": "",
            "business_number": "",
            "phone": "",
            "fax": "",
            "manager": "",
            "explainer": "",
        },
    }


def set_font(run, size: float = 8.5, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold


def format_paragraph(paragraph, size: float = 8.5, bold: bool = False, align=None, after: float = 0) -> None:
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    for run in paragraph.runs:
        set_font(run, size=size, bold=bold)


def set_cell_margins(cell, top: int = 60, start: int = 80, bottom: int = 60, end: int = 80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + side
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, size: float = 8.0, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    lines = str(text).splitlines() or [""]
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color: str = "000000", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(text)
    set_font(run, size=9.5, bold=True)


def add_two_column_header(doc: Document, data: dict[str, Any]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Mm(116)
    table.columns[1].width = Mm(64)
    set_cell_text(
        table.cell(0, 0),
        f"{data['metadata']['target_year']}　{data['metadata']['target_month']}月\n利用者氏名：{data['metadata']['user_name']}",
        size=9,
    )
    approvals = (
        "管理者　　　　　　サービス管理責任者\n"
        f"{data['metadata']['manager_approval']}　　　　　　{data['metadata']['service_manager_approval']}\n"
        f"作成日：{data['metadata']['created_date']}\n"
        f"前回作成日：{data['metadata']['previous_created_date']}\n"
        f"計画作成者：{data['metadata']['author']}"
    )
    set_cell_text(table.cell(0, 1), approvals, size=7.2)
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:val"), "nil")
            borders.append(node)
        tc_pr.append(borders)


def add_wishes(doc: Document, data: dict[str, Any]) -> None:
    add_heading(doc, "1. 利用者及び家族の要望")
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_cell_text(table.cell(0, 0), f"(1) 本人\n{data['wishes']['person']}")
    set_cell_text(table.cell(1, 0), f"(2) 家族\n{data['wishes']['family']}")


def add_goals(doc: Document, data: dict[str, Any]) -> None:
    add_heading(doc, "2. 支援目標と課題")
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    long_goal = data["goals"]["long_term"]
    short_goal = data["goals"]["short_term"]
    set_cell_text(table.cell(0, 0), f"(1) 長期　{long_goal['start_date']} ～ {long_goal['end_date']}\n{long_goal['text']}")
    set_cell_text(table.cell(1, 0), f"(2) 短期　{short_goal['start_date']} ～ {short_goal['end_date']}\n{short_goal['text']}")


def add_schedule(doc: Document, data: dict[str, Any]) -> None:
    add_heading(doc, f"4. 週間予定　{data['weekly_schedule']['legend']}")
    days = ["月", "火", "水", "木", "金", "土"]
    keys = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday"]
    rows = data["weekly_schedule"]["rows"]
    table = doc.add_table(rows=1 + len(rows), cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    headers = [""] + days
    widths = [Mm(34)] + [Mm(24.3)] * 6
    for index, header in enumerate(headers):
        table.columns[index].width = widths[index]
        set_cell_text(table.cell(0, index), header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.cell(0, index), "E7E6E6")
    set_repeat_table_header(table.rows[0])
    for row_index, item in enumerate(rows, start=1):
        set_cell_text(table.cell(row_index, 0), item["label"], size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        for col_index, key in enumerate(keys, start=1):
            set_cell_text(table.cell(row_index, col_index), item[key], size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_support_items(doc: Document, data: dict[str, Any]) -> None:
    add_heading(doc, "5. 具体的な課題及び支援計画等")
    intro = doc.add_paragraph("分野区分：①労働 ②訓練プログラム ③健康・医療 ④コミュニケーション ⑤日常生活 ⑥社会生活技能 ⑦社会参加 ⑧家族支援")
    format_paragraph(intro, size=6.8, after=1)
    table = doc.add_table(rows=1 + len(data["support_items"]), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    table.columns[0].width = Mm(16)
    table.columns[1].width = Mm(79)
    table.columns[2].width = Mm(85)
    headers = ["分野", "具体的な課題・到達目標", "支援内容（内容・方法・担当者等）"]
    for index, header in enumerate(headers):
        set_cell_text(table.cell(0, index), header, size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.cell(0, index), "E7E6E6")
    set_repeat_table_header(table.rows[0])
    for row_index, item in enumerate(data["support_items"], start=1):
        set_cant_split(table.rows[row_index])
        set_cell_text(table.cell(row_index, 0), "・".join(item["domains"]), size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        goal = f"{item['goal']}\n\n（{item['start_date']} ～ {item['end_date']}）"
        set_cell_text(table.cell(row_index, 1), goal, size=7.5)
        details = "\n".join(f"・{line}" if line else "" for line in item["support_details"])
        roles = "、".join(item["assigned_roles"])
        set_cell_text(table.cell(row_index, 2), f"{details}\n担当：{roles}", size=7.3)


def add_notes_and_agreement(doc: Document, data: dict[str, Any]) -> None:
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_cell_text(table.cell(0, 0), "備考", bold=True)
    set_cell_text(table.cell(1, 0), data["notes"] or "\n")

    agreement = data["agreement"]
    confirm = doc.add_table(rows=1, cols=2)
    confirm.alignment = WD_TABLE_ALIGNMENT.CENTER
    confirm.autofit = False
    set_table_borders(confirm)
    confirm.columns[0].width = Mm(82)
    confirm.columns[1].width = Mm(98)
    left = (
        "上記計画の内容について説明を受け同意しました。\n"
        f"日付：{agreement['explanation_date']}\n"
        f"ご本人氏名：{agreement['person_name']}\n"
        f"ご家族氏名：{agreement['family_name']}"
    )
    right = (
        "上記計画書に基づきサービスの説明を行い、内容に同意頂きましたので報告申し上げます。\n"
        f"同意日：{agreement['consent_date']}\n"
        f"確認日：{agreement['service_confirmation_date']}\n"
        f"確認先：{agreement['service_confirmation_name']}"
    )
    set_cell_text(confirm.cell(0, 0), left, size=7.2)
    set_cell_text(confirm.cell(0, 1), right, size=7.2)


def add_footer_info(doc: Document, data: dict[str, Any]) -> None:
    facility = data["facility"]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    table.columns[0].width = Mm(137)
    table.columns[1].width = Mm(43)
    left = (
        f"{facility['service_type']}　{facility['name']}　〒{facility['postal_code']}\n"
        f"住所：{facility['address']}　事業所No. {facility['business_number']}\n"
        f"Tel：{facility['phone']}　Fax：{facility['fax']}"
    )
    right = f"管理者：{facility['manager']}\n説明者：{facility['explainer']}"
    set_cell_text(table.cell(0, 0), left, size=6.8, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(0, 1), right, size=6.8)


def scrub_metadata(doc: Document) -> None:
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.keywords = ""
    props.subject = ""


def scrub_revision_ids(path: Path) -> None:
    """Remove Word revision-session identifiers from the generated package."""
    temp_path = path.with_suffix(".scrubbed.docx")
    with ZipFile(path, "r") as source, ZipFile(temp_path, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                try:
                    root = ET.fromstring(content)
                    for element in root.iter():
                        for attribute in list(element.attrib):
                            namespace, _, local_name = attribute[1:].partition("}") if attribute.startswith("{") else ("", "", attribute)
                            if namespace == WORD_NS and local_name.startswith("rsid"):
                                del element.attrib[attribute]
                    content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                except ET.ParseError:
                    pass
            target.writestr(info, content)
    temp_path.replace(path)


def build_document(data: dict[str, Any]) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.orientation = 0
    section.top_margin = Mm(10)
    section.bottom_margin = Mm(10)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(8.5)

    title_style = doc.styles["Title"]
    title_style.font.name = FONT_NAME
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    if title_style._element.pPr is not None:
        title_border = title_style._element.pPr.find(qn("w:pBdr"))
        if title_border is not None:
            title_style._element.pPr.remove(title_border)

    title = doc.add_paragraph()
    title.style = doc.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("個 別 支 援 計 画 書（案）")
    set_font(run, size=14, bold=True)

    add_two_column_header(doc, data)
    add_wishes(doc, data)
    add_goals(doc, data)
    add_schedule(doc, data)
    add_support_items(doc, data)
    add_notes_and_agreement(doc, data)
    add_footer_info(doc, data)
    scrub_metadata(doc)
    return doc


def main() -> int:
    args = parse_args()
    if args.output.suffix.lower() != ".docx":
        print("出力先は .docx で指定してください", file=sys.stderr)
        return 2

    try:
        schema = load_json(args.schema)
        data = blank_data() if args.blank else load_json(args.input)
        errors = validate(data, schema, schema)
        if errors:
            raise ValueError("入力検証に失敗しました:\n- " + "\n- ".join(errors))
        args.output.parent.mkdir(parents=True, exist_ok=True)
        build_document(data).save(args.output)
        scrub_revision_ids(args.output)
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        return 1

    print(f"Wordを生成しました: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
